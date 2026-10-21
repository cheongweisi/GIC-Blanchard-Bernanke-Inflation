"""Builds the Phase 5 Excel assumptions workbook.

Cell conventions (legend on the README sheet):
  light blue fill + red text   user-editable assumption (blank = documented default rule)
  light green fill             model-generated / transformed value (Excel formulas that mirror the engine; do not edit)
  grey fill                    historical realised data (History = information at the 2023Q2 origin; Realised = outturns after the origin,
                               used only for the conditional exercise B and for evaluation)
The engine (phase5/p5_forecast_engine.py) reads only hard-valued cells: History, Inputs_A/B, Settings, Aux_Models. The formula sheets
(Model_Inputs_A/B and the derived columns of History/Realised) exist so a user can see how each assumption is interpreted; the engine
recomputes every transformation itself and a test checks that the two agree.
"""
from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

USER_FILL = PatternFill("solid", fgColor="BDD7EE")
USER_FONT = Font(color="FF0000", name="Calibri", size=11)
DERIVED_FILL = PatternFill("solid", fgColor="E2EFDA")
REAL_FILL = PatternFill("solid", fgColor="D9D9D9")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
SUB_FONT = Font(italic=True, color="595959", name="Calibri", size=10)
TITLE_FONT = Font(bold=True, size=14, name="Calibri")
KEY_FONT = Font(color="A6A6A6", size=8, name="Calibri")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

HIST_START = "2019Q1"
ORIGIN = "2023Q2"
FIRST_FC = "2023Q3"
LAST_FC = "2029Q2"
MI_HIST_START = "2021Q1"
KEY_ROW, LABEL_ROW, UNIT_ROW, FIRST_DATA = 5, 6, 7, 8

RAW_COLS = [("brent", "Brent crude oil", "US$/bbl, quarterly avg"), ("cpi", "CPI-U all items (SA)", "index, quarter-end"),
            ("cpi_energy", "CPI energy (SA)", "index, quarter-end"), ("cpi_food", "CPI food (SA)", "index, quarter-end"),
            ("eci", "ECI wages & salaries, private (SA)", "index"), ("prod", "Nonfarm business productivity", "index"),
            ("gscpi", "GSCPI", "std devs, quarterly mean"), ("vu", "Job openings per unemployed (V/U)", "ratio"),
            ("cf1", "Cleveland Fed 1-year expected inflation", "% per year"), ("cf10", "Cleveland Fed 10-year expected inflation", "% per year")]
INPUT_COLS = [
    ("brent", "Brent crude oil", "US$/bbl, quarterly average", "Level. Blank = carried forward from the previous quarter."),
    ("cpie_qoq", "CPI energy: override of the oil bridge (optional)", "% change q/q, annualised",
     "Compounded annualised % change of the CPI energy index. If filled it replaces the oil-based path for that quarter. Blank = oil bridge."),
    ("cpif_qoq", "CPI food", "% change q/q, annualised",
     "Compounded annualised % change of the CPI food index. Blank = wage-linked: food prices rise with wages (relative food price unchanged)."),
    ("gscpi", "GSCPI (supply-chain pressure)", "std devs, quarterly mean", "Level. Blank = AR(1) mean reversion from the previous quarter."),
    ("vu", "Labour-market tightness V/U", "job openings per unemployed", "Level. Blank = carried forward from the previous quarter."),
    ("prod_qoq", "Nonfarm productivity growth", "% change q/q, annualised",
     "Compounded annualised % change. Blank = trend held (growth equal to 8 quarters earlier)."),
    ("af_gw", "Add-factor: wage equation", "pp, annualised", "Added to the wage-growth equation. Blank = 0."),
    ("af_gcpi", "Add-factor: CPI equation", "pp, annualised", "Added to the CPI-inflation equation. Blank = 0."),
    ("af_cf10", "Add-factor: 10y expectations eq.", "pp of % per year", "Added to the long-run expectations equation. Blank = 0."),
    ("af_cf1", "Add-factor: 1y expectations eq.", "pp of % per year", "Added to the short-run expectations equation. Blank = 0."),
]


def _quarters(a: str, b: str) -> list[str]:
    return [str(q) for q in pd.period_range(a, b, freq="Q")]


def _head(ws, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center", horizontal="center"), BORDER


def _title(ws, title: str, sub: str):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = sub
    ws["A2"].font = SUB_FONT
    ws["A2"].alignment = Alignment(wrap_text=False)


def _fmt(cell, fill=None, font=None, numfmt=None):
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if numfmt:
        cell.number_format = numfmt
    cell.border = BORDER


# ---------------------------------------------------------------------------------------------------------------------
def build_workbook(path, hist: pd.DataFrame, realised: pd.DataFrame, aux: dict, coef: pd.DataFrame, inputs_b: pd.DataFrame | None,
                   coef_file: str, meta: dict, forecast_end: str = "2026Q2", inputs_a: pd.DataFrame | None = None) -> None:
    """hist / realised: raw levels by key (index = Period), hist through the origin, realised = outturns after it (to 2026Q2).
    inputs_b: pre-fill of Inputs_B (index = Period, columns = INPUT keys); inputs_a: optional pre-fill of Inputs_A (default: blank)."""
    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "README"
    ws_set = wb.create_sheet("Settings")
    ws_ia, ws_ib = wb.create_sheet("Inputs_A"), wb.create_sheet("Inputs_B")
    ws_ma, ws_mb = wb.create_sheet("Model_Inputs_A"), wb.create_sheet("Model_Inputs_B")
    ws_var = wb.create_sheet("Variables")
    ws_h, ws_r = wb.create_sheet("History"), wb.create_sheet("Realised")
    ws_aux, ws_c = wb.create_sheet("Aux_Models"), wb.create_sheet("Coefficients")

    hq = _quarters(HIST_START, ORIGIN)
    fq = _quarters(FIRST_FC, LAST_FC)
    rq = _quarters(FIRST_FC, "2026Q2")
    h_row = {q: FIRST_DATA + i for i, q in enumerate(hq)}
    r_row = {q: FIRST_DATA + i for i, q in enumerate(rq)}

    _sheet_settings(ws_set, forecast_end, coef_file, meta)
    _sheet_aux(ws_aux, aux)
    _sheet_history(ws_h, hist, hq)
    _sheet_realised(ws_r, realised, rq, hq, h_row, meta)
    _sheet_inputs(ws_ia, "A", fq, inputs_a, rq, "Exercise A: pseudo-real-time assumptions (information available at 2023Q2)",
                  "Leave cells blank to use the documented default rule (see header notes and the Variables sheet). Everything blank = the pure "
                  "pseudo-real-time baseline: oil flat at its 2023Q2 level, V/U flat, GSCPI mean-reverting, trend productivity, food with wages.")
    _sheet_inputs(ws_ib, "B", fq, inputs_b, rq, "Exercise B: conditional assumptions (pre-filled with realised outturns through 2026Q2)",
                  "Pre-filled from the Realised sheet (historical realised values, for conditional analysis only). Overwrite any cell to run a "
                  "conditional scenario; clear a cell to fall back to the default rule. CPI-energy override takes precedence over Brent.")
    for ws, sc in ((ws_ma, "A"), (ws_mb, "B")):
        _sheet_model_inputs(ws, sc, fq, h_row, hq)
    _sheet_variables(ws_var, aux)
    _sheet_coef(ws_c, coef, coef_file)
    _sheet_readme(ws_readme, meta)
    wb.save(path)


# ---------------------------------------------------------------------------------------------------------------------
def _keyed_header(ws, keys, labels, units, ncols):
    ws.cell(KEY_ROW, 1, "key").font = KEY_FONT
    for j, k in enumerate(keys, start=1):
        if j > 1 and k:
            ws.cell(KEY_ROW, j, k).font = KEY_FONT
    for j, (lab, un) in enumerate(zip(labels, units), start=1):
        ws.cell(LABEL_ROW, j, lab)
        ws.cell(UNIT_ROW, j, un)
    _head(ws, LABEL_ROW, ncols)
    for c in range(1, ncols + 1):
        u = ws.cell(UNIT_ROW, c)
        u.font = SUB_FONT
        u.alignment = Alignment(wrap_text=True, horizontal="center", vertical="top")
        u.border = BORDER
    ws.row_dimensions[LABEL_ROW].height = 48
    ws.row_dimensions[UNIT_ROW].height = 30


def _sheet_settings(ws, forecast_end, coef_file, meta):
    _title(ws, "Settings", "The engine reads this sheet. Light-blue/red cells may be edited.")
    ws.cell(KEY_ROW, 1, "key").font = KEY_FONT
    ws.cell(LABEL_ROW, 1, "Setting")
    ws.cell(LABEL_ROW, 2, "Value")
    ws.cell(LABEL_ROW, 3, "Description")
    _head(ws, LABEL_ROW, 3)
    rows = [
        ("origin", ORIGIN, False, "Forecast origin: last quarter of data used for coefficients and initial conditions (fixed)."),
        ("forecast_end", forecast_end, True, "Last forecast quarter. Default 2026Q2 (end of the realised sample). Can be extended up to 2029Q2 "
                                              "(+3 years); the Inputs sheets already hold rows to 2029Q2."),
        ("coef_file", coef_file, False, "Estimated Model B* coefficients (1990Q1-2023Q2), path relative to this workbook. Never re-estimated by the engine."),
        ("model", "Model B* (four-equation SVAR-X)", False, "Authors' structure; GSCPI 2-quarter mean as the shortage term; Q2/Q3-2020 wage dummies (zero in forecasts)."),
        ("estimation_sample", meta.get("estimation_sample", "1990Q1-2023Q2"), False, "Estimation window of every equation."),
        ("data_vintage", meta.get("data_vintage", ""), False, "Data used for estimation and history."),
    ]
    # key rows must start right after the key row; keep order: key row holds the header 'key'
    for i, (k, v, user, desc) in enumerate(rows, start=KEY_ROW + 1):
        ws.cell(i, 1, k)
        c = ws.cell(i, 2, v)
        _fmt(c, USER_FILL if user else REAL_FILL, USER_FONT if user else None)
        ws.cell(i, 3, desc).alignment = WRAP
        ws.cell(i, 1).border = BORDER
        ws.cell(i, 3).border = BORDER
    dv = DataValidation(type="list", formula1='"' + ",".join(_quarters("2023Q3", LAST_FC)) + '"', allow_blank=False,
                        showErrorMessage=True, errorTitle="forecast_end", error=f"Choose a quarter from 2023Q3 to {LAST_FC}")
    ws.add_data_validation(dv)
    dv.add("B7")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 110
    assert ws["A7"].value == "forecast_end"


def _sheet_aux(ws, aux):
    _title(ws, "Auxiliary models (estimated on data through 2023Q2)",
           "Not part of Model B*'s four equations: they turn user-friendly assumptions into model inputs. Do not edit; re-estimate in Python instead.")
    ws.cell(KEY_ROW, 1, "key").font = KEY_FONT
    for j, t in enumerate(["Parameter", "Value", "Description"], start=1):
        ws.cell(LABEL_ROW, j, t)
    _head(ws, LABEL_ROW, 3)
    rows = [
        ("bridge_const", aux["bridge_const"], "Oil bridge constant: drift of CPI energy growth unrelated to oil (400*dln, % points)."),
        ("bridge_b0", aux["bridge_b0"], "Oil bridge pass-through: CPI energy growth = const + b0 * Brent growth (both 400*dln, quarterly)."),
        ("gscpi_ar_const", aux["gscpi_ar_const"], "GSCPI AR(1) constant (std devs)."),
        ("gscpi_ar_rho", aux["gscpi_ar_rho"], "GSCPI AR(1) persistence."),
        ("bridge_r2", aux["bridge_r2"], f"R-squared of the oil bridge ({aux['bridge_sample']}, n = {aux['bridge_n']}); residual s.d. {aux['bridge_resid_sd']:.1f} pp."),
        ("bridge_b0_se_ols", aux["bridge_b0_se_ols"], "OLS standard error of b0."),
        ("gscpi_ar_mean", aux["gscpi_ar_mean"], f"Implied long-run mean of GSCPI ({aux['gscpi_ar_sample']})."),
    ]
    for i, (k, v, d) in enumerate(rows, start=KEY_ROW + 1):
        ws.cell(i, 1, k).border = BORDER
        _fmt(ws.cell(i, 2, v), REAL_FILL, None, "0.0000")
        ws.cell(i, 3, d).border = BORDER
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 120
    assert ws["A6"].value == "bridge_const" and ws["A9"].value == "gscpi_ar_rho"


def _sheet_history(ws, hist, hq):
    _title(ws, "History: information available at the 2023Q2 origin (raw levels)",
           "Grey = historical data. Light green columns are formulas showing how each model variable is built from the levels. "
           "The engine uses the grey levels only and recomputes the model variables itself.")
    keys = ["quarter"] + [k for k, _, _ in RAW_COLS]
    der = [("gcpi", "CPI inflation gcpi", "400*dln, %"), ("gw", "Wage growth gw", "400*dln, %"), ("gE", "CPI energy growth", "400*dln, %"),
           ("gF", "CPI food growth", "400*dln, %"), ("grpe", "Relative energy price growth grpe", "400*dln(CPIE/ECI)"),
           ("grpf", "Relative food price growth grpf", "400*dln(CPIF/ECI)"), ("gpty", "Productivity growth", "400*dln, %"),
           ("magpty", "Trend productivity magpty", "8-qtr mean of gpty"), ("shortagema2", "Shortage term", "2-qtr mean of GSCPI"),
           ("diffcpicf", "Catch-up term", "4-qtr mean gcpi - cf1(t-4)")]
    labels = ["Quarter"] + [a for _, a, _ in RAW_COLS] + [a for _, a, _ in der]
    units = ["quarter"] + [u for _, _, u in RAW_COLS] + [u for _, _, u in der]
    ncols = len(labels)
    _keyed_header(ws, keys + [None] * len(der), labels, units, ncols)
    for i, q in enumerate(hq):
        r = FIRST_DATA + i
        ws.cell(r, 1, q).border = BORDER
        for j, (k, _, _) in enumerate(RAW_COLS, start=2):
            _fmt(ws.cell(r, j, float(hist.loc[pd.Period(q, "Q"), k])), REAL_FILL, None, "0.0000" if k in ("cpi", "cpi_energy", "cpi_food", "eci", "prod") else "0.000")
        c0 = 2 + len(RAW_COLS)
        col = {k: L(2 + i2) for i2, (k, _, _) in enumerate(RAW_COLS)}
        dc = {k: L(c0 + i2) for i2, (k, _, _) in enumerate(der)}
        f = {}
        if i >= 1:
            p = r - 1
            f["gcpi"] = f"=400*LN({col['cpi']}{r}/{col['cpi']}{p})"
            f["gw"] = f"=400*LN({col['eci']}{r}/{col['eci']}{p})"
            f["gE"] = f"=400*LN({col['cpi_energy']}{r}/{col['cpi_energy']}{p})"
            f["gF"] = f"=400*LN({col['cpi_food']}{r}/{col['cpi_food']}{p})"
            f["grpe"] = f"=400*LN(({col['cpi_energy']}{r}/{col['eci']}{r})/({col['cpi_energy']}{p}/{col['eci']}{p}))"
            f["grpf"] = f"=400*LN(({col['cpi_food']}{r}/{col['eci']}{r})/({col['cpi_food']}{p}/{col['eci']}{p}))"
            f["gpty"] = f"=400*LN({col['prod']}{r}/{col['prod']}{p})"
            f["shortagema2"] = f"=AVERAGE({col['gscpi']}{p}:{col['gscpi']}{r})"
        if i >= 8:
            f["magpty"] = f"=AVERAGE({dc['gpty']}{r - 7}:{dc['gpty']}{r})"
        if i >= 4:
            f["diffcpicf"] = f"=AVERAGE({dc['gcpi']}{r - 3}:{dc['gcpi']}{r})-{col['cf1']}{r - 4}"
        for k, _, _ in der:
            c = ws[f"{dc[k]}{r}"]
            if k in f:
                c.value = f[k]
            _fmt(c, DERIVED_FILL, None, "0.00")
    ws.freeze_panes = ws.cell(FIRST_DATA, 2)
    ws.column_dimensions["A"].width = 10
    for j in range(2, ncols + 1):
        ws.column_dimensions[L(j)].width = 14


def _sheet_realised(ws, realised, rq, hq, h_row, meta):
    _title(ws, "Realised: outturns after the origin (2023Q3-2026Q2)",
           "HISTORICAL REALISED VALUES, FOR CONDITIONAL SCENARIO ANALYSIS AND EVALUATION ONLY. Exercise A never reads this sheet.")
    keys = ["quarter"] + [k for k, _, _ in RAW_COLS]
    der = [("gcpi", "CPI inflation (target)", "400*dln, % q/q annualised"), ("cpi_yoy", "CPI inflation y/y", "%"),
           ("cpie_qoq", "CPI energy", "% q/q, annualised (Inputs_B unit)"), ("cpif_qoq", "CPI food", "% q/q, annualised (Inputs_B unit)"),
           ("prod_qoq", "Productivity growth", "% q/q, annualised (Inputs_B unit)")]
    labels = ["Quarter"] + [a for _, a, _ in RAW_COLS] + [a for _, a, _ in der] + ["Note"]
    units = ["quarter"] + [u for _, _, u in RAW_COLS] + [u for _, _, u in der] + [""]
    ncols = len(labels)
    _keyed_header(ws, keys + [None] * (len(der) + 1), labels, units, ncols)
    col = {k: L(2 + i2) for i2, (k, _, _) in enumerate(RAW_COLS)}
    c0 = 2 + len(RAW_COLS)
    dc = {k: L(c0 + i2) for i2, (k, _, _) in enumerate(der)}
    last_h = h_row[ORIGIN]
    flags = meta.get("realised_flags", {})
    for i, q in enumerate(rq):
        r = FIRST_DATA + i
        ws.cell(r, 1, q).border = BORDER
        for j, (k, _, _) in enumerate(RAW_COLS, start=2):
            _fmt(ws.cell(r, j, float(realised.loc[pd.Period(q, "Q"), k])), REAL_FILL, None, "0.0000" if k in ("cpi", "cpi_energy", "cpi_food", "eci", "prod") else "0.000")

        def prev(k, lag):
            # reference to the level `lag` quarters before row r: Realised sheet if available, else History
            idx = i - lag
            if idx >= 0:
                return f"{col[k]}{FIRST_DATA + idx}"
            hq_i = len(hq) - 1 + (idx + 1)
            return f"History!{col[k]}{h_row[hq[hq_i]]}"
        f = {
            "gcpi": f"=400*LN({col['cpi']}{r}/{prev('cpi', 1)})",
            "cpi_yoy": f"=100*({col['cpi']}{r}/{prev('cpi', 4)}-1)",
            "cpie_qoq": f"=100*(({col['cpi_energy']}{r}/{prev('cpi_energy', 1)})^4-1)",
            "cpif_qoq": f"=100*(({col['cpi_food']}{r}/{prev('cpi_food', 1)})^4-1)",
            "prod_qoq": f"=100*(({col['prod']}{r}/{prev('prod', 1)})^4-1)",
        }
        for k, _, _ in der:
            _fmt(ws[f"{dc[k]}{r}"], DERIVED_FILL, None, "0.00")
            ws[f"{dc[k]}{r}"].value = f[k]
        note = ws.cell(r, ncols)
        note.border = BORDER
        if q in flags.get("vu_bridged", []):
            note.value = "V/U bridged: October 2025 unemployment was never published (Phase 3 supplement); flagged"
    ws.freeze_panes = ws.cell(FIRST_DATA, 2)
    ws.column_dimensions["A"].width = 10
    for j in range(2, ncols):
        ws.column_dimensions[L(j)].width = 14
    ws.column_dimensions[L(ncols)].width = 70


def _sheet_inputs(ws, sc, fq, prefill, rq, title, sub):
    _title(ws, title, sub)
    keys = ["quarter"] + [None] + [k for k, *_ in INPUT_COLS] + [None]
    labels = ["Quarter", "Horizon status"] + [a for _, a, _, _ in INPUT_COLS] + ["Pre-fill source"]
    units = ["", "per Settings!forecast_end"] + [u for _, _, u, _ in INPUT_COLS] + [""]
    ncols = len(labels)
    _keyed_header(ws, keys, labels, units, ncols)
    for j, (k, lab, un, tip) in enumerate(INPUT_COLS, start=3):
        c = ws.cell(LABEL_ROW, j)
        c.comment = Comment(f"{lab} [{un}]\n{tip}", "Phase 5")
        c.comment.width, c.comment.height = 320, 120
    ws.cell(UNIT_ROW + 0, 1).value = ""
    for i, q in enumerate(fq):
        r = FIRST_DATA + i
        ws.cell(r, 1, q).border = BORDER
        _fmt(ws.cell(r, 2, f'=IF(A{r}<=Settings!$B$7,"in horizon","beyond horizon (ignored)")'), DERIVED_FILL)
        for j, (k, *_) in enumerate(INPUT_COLS, start=3):
            c = ws.cell(r, j)
            _fmt(c, USER_FILL, USER_FONT, "0.00" if not k.startswith("af_") else "0.00")
            if prefill is not None and pd.Period(q, "Q") in prefill.index:
                v = prefill.loc[pd.Period(q, "Q"), k]
                if pd.notna(v):
                    c.value = float(v)
        src = ws.cell(r, ncols)
        src.border = BORDER
        src.value = ("Realised outturn (Realised sheet); overwrite to run a conditional scenario"
                     if (prefill is not None and pd.Period(q, "Q") in prefill.index and prefill.loc[pd.Period(q, "Q")].notna().any())
                     else "blank: default rule")
    last = FIRST_DATA + len(fq) - 1
    dv_pos = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True, showErrorMessage=True,
                            errorTitle="Invalid", error="Must be a positive number")
    dv_pct = DataValidation(type="decimal", operator="greaterThan", formula1="-100", allow_blank=True, showErrorMessage=True,
                            errorTitle="Invalid", error="A % change must be greater than -100")
    dv_vu = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True,
                           errorTitle="Invalid", error="V/U cannot be negative")
    for dv in (dv_pos, dv_pct, dv_vu):
        ws.add_data_validation(dv)
    dv_pos.add(f"C{FIRST_DATA}:C{last}")
    for colx in ("D", "E", "H"):
        dv_pct.add(f"{colx}{FIRST_DATA}:{colx}{last}")
    dv_vu.add(f"G{FIRST_DATA}:G{last}")
    ws.freeze_panes = ws.cell(FIRST_DATA, 3)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 20
    for j in range(3, ncols):
        ws.column_dimensions[L(j)].width = 17
    ws.column_dimensions[L(ncols)].width = 58
    ws.cell(UNIT_ROW + 1, 1)  # spacing


def _sheet_model_inputs(ws, sc, fq, h_row, hq):
    _title(ws, f"Model inputs actually used (Exercise {sc}): transformed assumptions",
           "Light green = formulas. This is how the engine interprets Inputs_" + sc + " (the engine recomputes the same transformations in Python; "
           "a test confirms they agree). Relative prices grpe = gE - gw and grpf = gF - gw need the simulated wage growth, so they are produced by the engine "
           "(see the engine's Model_Inputs_Engine sheet).")
    cols = [("quarter", "Quarter", ""), ("status", "Status", ""),
            ("brent", "Brent used", "US$/bbl"), ("gbrent", "Brent growth", "400*dln"), ("gEb", "CPI energy growth from oil bridge", "400*dln"),
            ("gE", "CPI energy growth USED (model gE)", "400*dln"), ("gF", "CPI food growth USED (model gF)", "400*dln"),
            ("gscpi", "GSCPI used", "std devs"), ("sma2", "Shortage term (2-qtr mean)", "std devs"), ("vu", "V/U used", "ratio"),
            ("gpty", "Productivity growth used", "400*dln"), ("magpty", "Trend productivity magpty (8-qtr mean)", "400*dln"),
            ("af_gw", "Add-factor gw", "pp"), ("af_gcpi", "Add-factor gcpi", "pp"), ("af_cf10", "Add-factor cf10", "pp"), ("af_cf1", "Add-factor cf1", "pp"),
            ("s_brent", "Brent: source", ""), ("s_gE", "CPI energy: source", ""), ("s_gF", "CPI food: source", ""),
            ("s_gscpi", "GSCPI: source", ""), ("s_vu", "V/U: source", ""), ("s_prod", "Productivity: source", "")]
    c = {k: L(i) for i, (k, _, _) in enumerate(cols, start=1)}
    for j, (_, lab, un) in enumerate(cols, start=1):
        ws.cell(LABEL_ROW, j, lab)
        ws.cell(UNIT_ROW, j, un)
    _head(ws, LABEL_ROW, len(cols))
    for j in range(1, len(cols) + 1):
        u = ws.cell(UNIT_ROW, j)
        u.font, u.alignment, u.border = SUB_FONT, Alignment(horizontal="center"), BORDER
    ws.row_dimensions[LABEL_ROW].height = 60
    hsel = [q for q in hq if q >= MI_HIST_START]
    n_h = len(hsel)
    inp = f"Inputs_{sc}"
    hc = {k: L(2 + i) for i, (k, _, _) in enumerate(RAW_COLS)}          # History raw columns
    hd0 = 2 + len(RAW_COLS)
    hd = {k: L(hd0 + i) for i, k in enumerate(["gcpi", "gw", "gE", "gF", "grpe", "grpf", "gpty", "magpty", "shortagema2", "diffcpicf"])}
    for i, q in enumerate(hsel):
        r = FIRST_DATA + i
        hr = h_row[q]
        vals = {"quarter": q, "status": "History", "brent": f"=History!{hc['brent']}{hr}", "gscpi": f"=History!{hc['gscpi']}{hr}",
                "sma2": f"=History!{hd['shortagema2']}{hr}", "vu": f"=History!{hc['vu']}{hr}", "gE": f"=History!{hd['gE']}{hr}",
                "gF": f"=History!{hd['gF']}{hr}", "gpty": f"=History!{hd['gpty']}{hr}", "magpty": f"=History!{hd['magpty']}{hr}"}
        if i >= 1:
            vals["gbrent"] = f"=400*LN({c['brent']}{r}/{c['brent']}{r - 1})"
        for k, v in vals.items():
            cell = ws[f"{c[k]}{r}"]
            cell.value = v
            _fmt(cell, REAL_FILL if k in ("quarter", "status") else DERIVED_FILL, None, "0.00")
    first_f = FIRST_DATA + n_h
    for i, q in enumerate(fq):
        r = first_f + i
        ri = FIRST_DATA + i                      # Inputs row
        p = r - 1
        ref = lambda col: f"{inp}!{col}{ri}"      # noqa: E731
        f = {
            "quarter": q,
            "status": f'=IF({c["quarter"]}{r}<=Settings!$B$7,"Forecast","Beyond horizon")',
            "brent": f"=IF(ISBLANK({ref('C')}),{c['brent']}{p},{ref('C')})",
            "gbrent": f"=400*LN({c['brent']}{r}/{c['brent']}{p})",
            "gEb": f"=Aux_Models!$B$6+Aux_Models!$B$7*{c['gbrent']}{r}",
            "gE": f"=IF(ISBLANK({ref('D')}),{c['gEb']}{r},100*LN(1+{ref('D')}/100))",
            "gF": f'=IF(ISBLANK({ref("E")}),"wage-linked",100*LN(1+{ref("E")}/100))',
            "gscpi": f"=IF(ISBLANK({ref('F')}),Aux_Models!$B$8+Aux_Models!$B$9*{c['gscpi']}{p},{ref('F')})",
            "sma2": f"=AVERAGE({c['gscpi']}{p}:{c['gscpi']}{r})",
            "vu": f"=IF(ISBLANK({ref('G')}),{c['vu']}{p},{ref('G')})",
            "gpty": f"=IF(ISBLANK({ref('H')}),{c['gpty']}{r - 8},100*LN(1+{ref('H')}/100))",
            "magpty": f"=AVERAGE({c['gpty']}{r - 7}:{c['gpty']}{r})",
            "af_gw": f"=IF(ISBLANK({ref('I')}),0,{ref('I')})", "af_gcpi": f"=IF(ISBLANK({ref('J')}),0,{ref('J')})",
            "af_cf10": f"=IF(ISBLANK({ref('K')}),0,{ref('K')})", "af_cf1": f"=IF(ISBLANK({ref('L')}),0,{ref('L')})",
            "s_brent": f'=IF(ISBLANK({ref("C")}),"default: carried forward","user")',
            "s_gE": f'=IF(NOT(ISBLANK({ref("D")})),"user (CPI energy override)",IF(ISBLANK({ref("C")}),"oil bridge, Brent default","oil bridge, Brent user"))',
            "s_gF": f'=IF(ISBLANK({ref("E")}),"default: wage-linked","user")',
            "s_gscpi": f'=IF(ISBLANK({ref("F")}),"default: AR(1)","user")',
            "s_vu": f'=IF(ISBLANK({ref("G")}),"default: carried forward","user")',
            "s_prod": f'=IF(ISBLANK({ref("H")}),"default: trend held","user")',
        }
        for k, v in f.items():
            cell = ws[f"{c[k]}{r}"]
            cell.value = v
            _fmt(cell, REAL_FILL if k == "quarter" else DERIVED_FILL, None, "0.00")
    ws.freeze_panes = ws.cell(FIRST_DATA, 3)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 16
    for j in range(3, len(cols) + 1):
        ws.column_dimensions[L(j)].width = 15 if j < 17 else 27


def _sheet_variables(ws, aux):
    _title(ws, "Variables: every exogenous input of Model B*, its units, transformation and treatment",
           "Category: USER = user-specified assumption (blue/red cells); MODEL = model-generated projection; REALISED = historical realised value (conditional analysis only).")
    hdr = ["Variable (Model B* name)", "Role in Model B*", "Category", "User unit (workbook)", "Model unit and scaling", "Transformation (engine = Model_Inputs sheet)",
           "Historical-to-forecast transition", "If left blank (default rule)", "Equation(s) and lags"]
    for j, h in enumerate(hdr, start=1):
        ws.cell(5, j, h)
    _head(ws, 5, len(hdr))
    b0, bc = aux["bridge_b0"], aux["bridge_const"]
    rho, gc = aux["gscpi_ar_rho"], aux["gscpi_ar_const"]
    R = [
        ("Brent crude oil", "Driver of CPI energy through the oil bridge (not itself a Model B* regressor)", "USER (Exercise B: pre-filled REALISED)",
         "US$ per barrel, quarterly average level", "gBrent = 400*dln(Brent), % points annualised",
         f"gBrent_t = 400*LN(P_t/P_(t-1)); CPI-energy growth gE_t = {bc:.3f} + {b0:.3f}*gBrent_t (bridge estimated 1990Q2-2023Q2). Entering 80 uses ln(80/previous level), never 80.",
         "First forecast quarter uses the 2023Q2 average (History) as P_(t-1).", "Carried forward from the previous quarter (first quarter: 2023Q2 average).",
         "Enters the CPI equation via grpe (lags 0-4)"),
        ("CPI energy (override)", "grpe = relative price of energy (CPI energy / ECI wages), growth", "USER (optional; Exercise B: pre-filled REALISED)",
         "% change q/q, annualised (compounded)", "gE = 400*dln(CPI energy) = 100*ln(1 + x/100)",
         "gE_t = 100*LN(1 + x/100); then grpe_t = gE_t - gw_t, computed every quarter from simulated wage growth (wages are endogenous, so the relative price is an OUTPUT).",
         "Quarter-end index convention; history grpe from History.", "Oil bridge applied to the Brent path.", "Price equation: grpe lags 0-4"),
        ("CPI food", "grpf = relative price of food, growth", "USER (Exercise B: pre-filled REALISED)",
         "% change q/q, annualised (compounded)", "gF = 400*dln(CPI food) = 100*ln(1 + x/100)",
         "gF_t = 100*LN(1 + x/100); grpf_t = gF_t - gw_t (derived each quarter).", "History grpf from History.",
         "Wage-linked: nominal food grows with simulated wages, i.e. grpf = 0 (the authors' neutral convention).", "Price equation: grpf lags 0-4"),
        ("GSCPI (supply-chain pressure)", "shortagema2 = shortage term of the price equation", "USER (Exercise B: pre-filled REALISED)",
         "Standard deviations, quarterly mean level", "level (std devs)", "shortagema2_t = (GSCPI_t + GSCPI_(t-1)) / 2 (two-quarter trailing mean)",
         "GSCPI_(t-1) for the first forecast quarter is the 2023Q2 value in History.",
         f"AR(1): GSCPI_t = {gc:.4f} + {rho:.3f}*GSCPI_(t-1) (estimated 1998Q1-2023Q2).", "Price equation: shortagema2 lag 0"),
        ("Labour-market tightness V/U", "vu", "USER (Exercise B: pre-filled REALISED; 2025Q4 bridged)", "Job openings per unemployed worker (ratio)", "level (ratio)",
         "none (used as entered)", "Lags 1-4 use History for the first forecast quarters: V/U affects wages with a one-quarter delay.",
         "Carried forward from the previous quarter (first quarter: 2023Q2 value).", "Wage equation: vu lags 1-4"),
        ("Productivity growth (nonfarm business, output per hour)", "magpty = trend productivity growth", "USER (Exercise B: pre-filled REALISED)",
         "% change q/q, annualised (compounded)", "gpty = 400*dln(productivity); magpty = 8-quarter trailing mean of gpty",
         "gpty_t = 100*LN(1 + x/100); magpty_t = AVERAGE(gpty_(t-7..t))",
         "The 8-quarter window is filled from History (2021Q3-2023Q2) and rolls forward.",
         "Trend held: gpty_t = gpty_(t-8), so magpty stays at its 2023Q2 value.", "Wage equation: magpty lag 1; price equation: magpty lag 0"),
        ("Equation add-factors (4)", "Intercept shifts / residual overrides", "USER", "Percentage points (same unit as the dependent variable)",
         "gw, gcpi: annualised % points; cf10, cf1: % per year", "added to the right-hand side of the equation in the quarter",
         "n.a.", "Zero.", "One per equation"),
        ("COVID dummies D2020Q2, D2020Q3", "Wage-equation controls estimated in sample", "MODEL (fixed at zero)", "-", "0/1", "fixed 0 in every forecast quarter",
         "Both quarters are inside the estimation sample; no pandemic quarter is assumed in the horizon.", "Zero (not editable).", "Wage equation, lag 0"),
        ("Catch-up term diffcpicf", "Endogenous function of CPI inflation and 1y expectations", "MODEL", "-", "pp: mean(gcpi_t..t-3) - cf1_(t-4)",
         "computed every quarter after gcpi is solved", "Uses History for the first four forecast quarters.", "n.a. (model-generated)", "Wage equation: lags 1-4"),
        ("gw, gcpi, cf10, cf1 (endogenous)", "The four equations' dependent variables", "MODEL", "-", "gw, gcpi: 400*dln; cf10, cf1: % per year",
         "solved recursively: gw -> grpe, grpf -> gcpi -> catch-up -> cf10 -> cf1", "Initial conditions = last four History quarters.", "n.a.", "-"),
        ("Realised data after 2023Q2", "Outturns", "REALISED", "see Realised sheet", "-", "Used only (i) to pre-fill Inputs_B and (ii) to compute forecast errors after the run.",
         "Never enters Exercise A, the coefficients or any initial condition.", "-", "-"),
    ]
    for i, row in enumerate(R, start=6):
        for j, v in enumerate(row, start=1):
            c = ws.cell(i, j, v)
            c.alignment, c.border = WRAP, BORDER
        cat = ws.cell(i, 3)
        cat.fill = USER_FILL if cat.value.startswith("USER") else (DERIVED_FILL if cat.value.startswith("MODEL") else REAL_FILL)
        ws.row_dimensions[i].height = 92
    for j, w in enumerate([26, 30, 24, 26, 30, 60, 42, 44, 28], start=1):
        ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = "B6"


def _sheet_coef(ws, coef, coef_file):
    _title(ws, "Estimated Model B* coefficients (1990Q1-2023Q2), display only",
           f"The engine reads {coef_file}; this sheet is a copy for reference and is cross-checked against the file at run time. Do not edit.")
    ws.cell(KEY_ROW, 1, "key").font = KEY_FONT
    for j, t in enumerate(["", "Equation", "Term (variable_Llag)", "Coefficient", "HAC s.e.", "HAC t"], start=1):
        ws.cell(LABEL_ROW, j, t)
    _head(ws, LABEL_ROW, 6)
    for i, r in enumerate(coef.itertuples(), start=FIRST_DATA):
        ws.cell(i, 2, r.eq)
        ws.cell(i, 3, r.term)
        _fmt(ws.cell(i, 4, float(r.coef)), REAL_FILL, None, "0.000000")
        _fmt(ws.cell(i, 5, float(r.se_hac)), REAL_FILL, None, "0.0000")
        _fmt(ws.cell(i, 6, float(r.t_hac)), REAL_FILL, None, "0.00")
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 24
    for c in "DEF":
        ws.column_dimensions[c].width = 14


def _sheet_readme(ws, meta):
    ws["A1"] = "Phase 5: Model B* inflation forecasting workbook (origin 2023Q2)"
    ws["A1"].font = TITLE_FONT
    lines = [
        ("What this is", "The editable assumptions for the Model B* forecasting engine. You enter assumptions in everyday units (Brent in US$/bbl, "
                         "annualised % changes, levels); the engine converts them into the exact variables the estimated equations use. Coefficients are fixed "
                         "(estimated on 1990Q1-2023Q2) and are never re-estimated when you change an assumption."),
        ("How to run", "Edit Inputs_A and/or Inputs_B (and Settings!forecast_end), save, then from the phase5 folder run:\n"
                       "python p5_forecast_engine.py --workbook Phase5_Assumptions.xlsx --scenario A --with-realised\n"
                       "python p5_forecast_engine.py --workbook Phase5_Assumptions.xlsx --scenario B --with-realised\n"
                       "Results are written to output/ (CSV and Excel with the transformed inputs the engine used and a log of every default rule applied)."),
        ("Exercise A", "Pseudo-real-time forecast: only information available at 2023Q2. All Inputs_A cells are blank, which applies the default rules "
                       "(oil flat at its 2023Q2 level, V/U flat, GSCPI AR(1), trend productivity, food with wages)."),
        ("Exercise B", "Conditional forecast: the realised paths of the drivers (pre-filled in Inputs_B from the Realised sheet) with the same fixed coefficients. "
                       "Wages and all four endogenous variables are still simulated."),
        ("Colour legend", "LIGHT BLUE fill + RED text = cell you may edit (user-specified assumption; blank = default rule)\n"
                          "LIGHT GREEN fill = model-generated or transformed value (formulas; do not edit)\n"
                          "GREY fill = historical realised data (History: known at the origin; Realised: outturns after it, conditional analysis only)"),
        ("Sheets", "Inputs_A / Inputs_B: your assumptions, quarterly 2023Q3-2029Q2.\nModel_Inputs_A / Model_Inputs_B: the transformed model inputs, so you can check "
                   "that an assumption was interpreted as intended (e.g. Brent 80 becomes a log change, not the number 80).\nVariables: units, transformations, "
                   "transition rules and default for every input.\nHistory / Realised: raw levels.\nSettings, Aux_Models, Coefficients: model parameters (fixed)."),
        ("Key transformations", "Brent: 400*ln(P_t/P_(t-1)), then CPI-energy growth = const + b0*that (oil bridge). CPI energy, CPI food, productivity: annualised % -> "
                                "100*ln(1+x/100) (= 400*ln of the quarterly ratio). GSCPI: 2-quarter trailing mean. V/U: level. Productivity trend: 8-quarter trailing mean. "
                                "Relative prices grpe, grpf = nominal growth minus simulated wage growth (derived every quarter)."),
        ("Extending the horizon", "Set Settings!forecast_end up to 2029Q2 (three years beyond 2026Q2). Quarters without inputs use the default rules."),
        ("Caveats", "Model B* is a reduced-form wage-price system estimated on revised data; oil, food, labour market, productivity and supply chains are exogenous "
                    "here and the main source of forecast error. See METHODOLOGY_NOTE.md for the assessment and limitations."),
        ("Data", meta.get("data_vintage", "")),
    ]
    for i, (k, v) in enumerate(lines, start=3):
        a, b = ws.cell(i, 1, k), ws.cell(i, 2, v)
        a.font = Font(bold=True)
        a.alignment = b.alignment = WRAP
        ws.row_dimensions[i].height = max(30, 15 * (len(v) // 100 + 1 + v.count("\n")))
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 130
    ws["A14"] = "Examples"
    ws["A14"].font = Font(bold=True)
    ws["B14"] = "Editable example cell:"
    ex = ws["C14"]
    ex.value = 80
    _fmt(ex, USER_FILL, USER_FONT)
    ws["D14"] = "model-generated example:"
    d = ws["E14"]
    d.value = 4.2
    _fmt(d, DERIVED_FILL)
    ws["F14"] = "historical realised example:"
    e = ws["G14"]
    e.value = 3.1
    _fmt(e, REAL_FILL)
