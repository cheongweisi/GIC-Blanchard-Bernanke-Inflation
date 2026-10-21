"""Phase 5: estimation through 2023Q2 uses no later data, the standalone engine is exact and leak-free, transformations are right,
the decomposition is additive, and Phases 1-4c are untouched."""
import hashlib
import shutil
import sys

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from bb_inflation.config import PROJECT_ROOT
from bb_inflation.phase5 import data as D
from bb_inflation.phase5 import estimate as E
from bb_inflation.phase5 import evaluate as EV
from bb_inflation.phase5 import pipeline as PL
from bb_inflation.scenarios import engine as SE
from bb_inflation.simulation.model import SimModel

P5 = PROJECT_ROOT / "phase5"
sys.path.insert(0, str(P5))
import p5_forecast_engine as eng  # noqa: E402

WB = P5 / "Phase5_Assumptions.xlsx"
COEFS = P5 / "coefficients" / "Bstar_coefficients_2023Q2.csv"
pytestmark = pytest.mark.skipif(not WB.exists(), reason="run phase5/run_phase5.py first")


@pytest.fixture(scope="module")
def p5():
    return D.load()


@pytest.fixture(scope="module")
def wd():
    return eng.read_workbook(WB)


@pytest.fixture(scope="module")
def coefs():
    return eng.read_coefficients(COEFS)


def test_existing_files_unchanged():
    bad = []
    for line in (P5 / "pre_phase5_manifest.sha256").read_text().splitlines():
        h, rel = line.split("  ", 1)
        f = PROJECT_ROOT / rel
        if not f.exists() or hashlib.sha256(f.read_bytes()).hexdigest() != h:
            bad.append(rel)
    assert not bad, bad[:5]


# ---------------------------------------------------------------- estimation: no use of data after 2023Q2
def test_estimation_ignores_everything_after_2023q2(p5):
    data = p5.model
    base = E.fit(data, "2023Q2")
    poisoned = data.copy()
    poisoned.loc["2023Q3":, :] = poisoned.loc["2023Q3":, :] * 3.7 + 11.0         # wreck every post-origin observation
    alt = E.fit(poisoned, "2023Q2")
    for eq in base:
        assert np.allclose(base[eq].coef.values, alt[eq].coef.values, atol=0, rtol=0)
        assert base[eq].sample[1] == pd.Period("2023Q2", "Q") and base[eq].nobs == 134


def test_data_truncated_at_origin_gives_same_dataset(p5):
    from bb_inflation.data import phase3_data as P3
    from bb_inflation.supplement import context as C
    d2 = C.model_data(P3.build(through="2023Q2").raw)
    a, b = d2.loc[:"2023Q2"], p5.model.loc[:"2023Q2"]
    assert np.nanmax(np.abs(a.values - b[a.columns].values)) < 1e-12


def test_coefficient_file_matches_fresh_estimate_and_restrictions(p5, coefs):
    res = E.fit(p5.model, "2023Q2")
    for eq, r in res.items():
        assert np.allclose(coefs[eq].reindex(r.coef.index).values, r.coef.values, atol=1e-12)
    eng.validate_coefficients(coefs)                                  # exact term set and sum-to-one
    bad = {k: v.copy() for k, v in coefs.items()}
    bad["gcpi"]["gw_L1"] += 0.01
    with pytest.raises(ValueError):
        eng.validate_coefficients(bad)
    extra = {k: v.copy() for k, v in coefs.items()}
    extra["cf10"]["vu_L1"] = 0.0
    with pytest.raises(ValueError):
        eng.validate_coefficients(extra)


def test_2026_fit_is_production_bstar(p5):
    st = pd.read_csv(PROJECT_ROOT / "outputs" / "phase4c" / "coef_gcpi.csv")
    stored = st.set_index("coefficient")["B*-2026: estimate"].dropna()
    new = E.fit(p5.model, "2026Q2")["gcpi"].coef
    assert np.allclose(stored.reindex(new.index).values, new.values, atol=1e-8)


def test_matches_phase3_model_a_frozen_at_2023q2(p5):
    from bb_inflation.extension.models import ModelSpec, fit
    a = fit(ModelSpec("A", "A", end="2023Q2"), p5.model, "2026Q2")
    new = E.fit(p5.model, "2023Q2")
    for eq in new:
        assert np.allclose(a[eq].coef.reindex(new[eq].coef.index).values, new[eq].coef.values, atol=1e-10)


# ---------------------------------------------------------------- engine mechanics
def test_derived_history_equals_estimation_dataset(wd, p5):
    d = eng.derive_history(wd.history)
    ref = p5.model
    for k in ("gcpi", "gw", "grpe", "grpf", "magpty", "shortagema2", "diffcpicf", "vu", "cf1", "cf10"):
        v = d[k].dropna().loc["2021Q2":]
        assert np.allclose(v.values, ref.loc[v.index, k].values, atol=1e-9), k


def test_matrix_form_equals_sequential_solution(wd, coefs):
    for sc in "AB":
        r = eng.run_scenario(wd, sc, coefs=coefs)
        assert r.checks["matrix_form_max_abs_diff"] < 1e-10


def test_engine_equals_phase4_simulator(wd, coefs):
    """Same recursion as the validated Phase 4 nominal-price engine (which reproduces the Phase 1-3 simulator to 1e-15)."""
    r = eng.run_scenario(wd, "B", coefs=coefs)
    d = eng.derive_history(wd.history)
    h = d.iloc[-4:]
    mi = r.inputs_used
    H = len(mi)
    cs = SE.CoefSet.from_models([SimModel({eq: c for eq, c in coefs.items()})])
    hist = {k: h[k].values for k in ("gw", "gcpi", "cf1", "cf10", "diffcpicf", "grpe", "grpf")}
    inp = SE.Inputs(hist=hist, exog={"vu": np.r_[h["vu"].values, mi["vu"].values], "magpty": np.r_[h["magpty"].values, mi["magpty"].values],
                                     "shortagema2": np.r_[h["shortagema2"].values, mi["shortagema2"].values]},
                    gE=np.r_[h["gE"].values, mi["gE"].values], gF=np.r_[h["gF"].values, mi["gF"].values])
    out = SE.simulate(cs, inp)
    assert np.allclose(out["gcpi"][0, 4:], r.forecast["gcpi"].values, atol=1e-10)
    assert np.allclose(out["cf1"][0, 4:], r.forecast["cf1"].values, atol=1e-10)


def test_exercise_b_with_all_residuals_reproduces_realised(wd, coefs, p5):
    q = pd.period_range("2023Q3", "2026Q2", freq="Q")
    res23 = E.fit(p5.model, "2023Q2")
    resid = EV.one_step_residuals(res23, p5.model, q)
    r = eng.run_scenario(wd, "B", coefs=coefs, extra_add={e: resid[e].values for e in resid})
    real = p5.model.loc[q]
    for k in ("gw", "gcpi", "cf1", "cf10", "grpe", "grpf"):
        assert np.allclose(r.forecast[k].values, real[k].values, atol=1e-8), k


# ---------------------------------------------------------------- transformations
def test_brent_80_is_converted_not_fed_in(wd, coefs):
    inp = wd.inputs["A"].copy()
    inp.loc[pd.Period("2023Q3", "Q"), "brent"] = 80.0
    r = eng.run_scenario(wd, "A", coefs=coefs, inputs_override=inp)
    mi = r.inputs_used
    last = float(wd.history["brent"].iloc[-1])
    assert mi.loc["2023Q3", "brent"] == 80.0
    assert mi.loc["2023Q3", "gbrent"] == pytest.approx(400 * np.log(80.0 / last))
    assert mi.loc["2023Q3", "gE"] == pytest.approx(wd.aux["bridge_const"] + wd.aux["bridge_b0"] * 400 * np.log(80.0 / last))
    assert abs(mi.loc["2023Q3", "gE"]) < 50                           # nothing like 80
    assert mi.loc["2023Q4", "gbrent"] == pytest.approx(0.0)           # level 80 carried forward: no further change


def test_percent_inputs_use_100_ln_1_plus_x(wd, coefs):
    inp = wd.inputs["A"].copy()
    inp.loc[pd.Period("2024Q1", "Q"), ["cpif_qoq", "cpie_qoq", "prod_qoq"]] = [6.0, -10.0, 2.5]
    mi = eng.run_scenario(wd, "A", coefs=coefs, inputs_override=inp).inputs_used
    assert mi.loc["2024Q1", "gF"] == pytest.approx(100 * np.log(1.06))
    assert mi.loc["2024Q1", "gE"] == pytest.approx(100 * np.log(0.90))
    assert mi.loc["2024Q1", "gpty"] == pytest.approx(100 * np.log(1.025))
    # compounding check: four quarters at 100*ln(1+x/100)/... annualised rate x -> 400*dln over four quarters sums to 400*ln(1+x/100)
    assert np.exp(4 * mi.loc["2024Q1", "gF"] / 400) == pytest.approx(1.06)
    assert eng.model_to_pct_ann(eng.pct_ann_to_model(7.3)) == pytest.approx(7.3)


def test_gscpi_ma2_vu_and_productivity_trend(wd, coefs):
    inp = wd.inputs["A"].copy()
    inp.loc[pd.Period("2023Q3", "Q"), "gscpi"] = 1.0
    inp.loc[pd.Period("2023Q4", "Q"), "gscpi"] = 3.0
    inp.loc[pd.Period("2023Q4", "Q"), "vu"] = 1.1
    r = eng.run_scenario(wd, "A", coefs=coefs, inputs_override=inp)
    mi = r.inputs_used
    g0 = float(wd.history["gscpi"].iloc[-1])
    assert mi.loc["2023Q3", "shortagema2"] == pytest.approx(0.5 * (1.0 + g0))
    assert mi.loc["2023Q4", "shortagema2"] == pytest.approx(2.0)
    assert mi.loc["2023Q4", "vu"] == 1.1 and mi.loc["2024Q1", "vu"] == 1.1          # a user value sticks
    d = eng.derive_history(wd.history)
    assert np.allclose(mi["magpty"].values, d["magpty"].iloc[-1])                    # default: trend held at the 2023Q2 value


def test_blank_inputs_use_documented_defaults_and_are_logged(wd, coefs):
    r = eng.run_scenario(wd, "A", coefs=coefs)
    assert set(r.log["variable"]) == {"brent", "cpi_energy", "cpi_food", "gscpi", "vu", "productivity"}
    mi = r.inputs_used
    assert (mi["brent"] == wd.history["brent"].iloc[-1]).all() and (mi["vu"] == wd.history["vu"].iloc[-1]).all()
    assert mi["gF"].isna().all() and np.allclose(r.forecast["grpf"].values, 0.0)          # wage-linked food: relative price unchanged
    ar = mi["gscpi"].iloc[0]
    assert ar == pytest.approx(wd.aux["gscpi_ar_const"] + wd.aux["gscpi_ar_rho"] * wd.history["gscpi"].iloc[-1])
    assert (mi[["af_gw", "af_gcpi", "af_cf10", "af_cf1"]] == 0).all().all()


def test_relative_prices_are_derived_from_simulated_wages(wd, coefs):
    r = eng.run_scenario(wd, "B", coefs=coefs)
    mi, f = r.inputs_used, r.forecast
    assert np.allclose(f["grpe"].values, mi["gE"].values - f["gw"].values)
    assert np.allclose(f["grpf"].values, mi["gF"].values - f["gw"].values)


# ---------------------------------------------------------------- leakage and workbook handling
def test_exercise_a_never_reads_realised(tmp_path, coefs):
    wb = load_workbook(WB)
    ws = wb["Realised"]
    for row in ws.iter_rows(min_row=8, max_row=19, min_col=2, max_col=11):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.value = c.value * 2.5 + 7
    f = tmp_path / "poisoned.xlsx"
    wb.save(f)
    a0 = eng.run_scenario(WB, "A", coefs=coefs).forecast
    a1 = eng.run_scenario(f, "A", coefs=coefs).forecast
    pd.testing.assert_frame_equal(a0[["gw", "gcpi", "cf1", "cf10", "cpi", "cpi_yoy"]], a1[["gw", "gcpi", "cf1", "cf10", "cpi", "cpi_yoy"]])


def test_history_ends_at_origin_and_realised_only_after(wd):
    assert wd.history.index[-1] == pd.Period("2023Q2", "Q")
    assert wd.realised.index[0] == pd.Period("2023Q3", "Q") and wd.realised.index[-1] == pd.Period("2026Q2", "Q")


def test_horizon_extension_to_2029q2_runs_with_defaults(wd, coefs):
    r = eng.run_scenario(wd, "B", coefs=coefs, horizon="2029Q2")
    assert len(r.forecast) == 24 and r.forecast["gcpi"].notna().all()
    first12 = eng.run_scenario(wd, "B", coefs=coefs, horizon="2026Q2").forecast
    assert np.allclose(first12["gcpi"].values, r.forecast["gcpi"].iloc[:12].values)      # extension does not alter earlier quarters
    with pytest.raises(ValueError):
        eng.run_scenario(wd, "B", coefs=coefs, horizon="2029Q3")


def test_bad_inputs_are_rejected(wd, coefs):
    inp = wd.inputs["A"].copy()
    inp.loc[pd.Period("2023Q3", "Q"), "brent"] = -5.0
    with pytest.raises(ValueError):
        eng.run_scenario(wd, "A", coefs=coefs, inputs_override=inp)
    inp = wd.inputs["A"].copy()
    inp.loc[pd.Period("2023Q3", "Q"), "cpif_qoq"] = -120.0
    with pytest.raises(ValueError):
        eng.run_scenario(wd, "A", coefs=coefs, inputs_override=inp)


def test_text_in_input_cell_is_rejected(tmp_path):
    wb = load_workbook(WB)
    wb["Inputs_A"]["C8"] = "eighty"
    f = tmp_path / "bad.xlsx"
    wb.save(f)
    with pytest.raises(eng.WorkbookError):
        eng.read_workbook(f)


def test_editing_assumptions_does_not_change_coefficients(wd, coefs):
    before = {k: v.copy() for k, v in coefs.items()}
    inp = wd.inputs["A"].copy()
    inp["brent"] = 120.0
    eng.run_scenario(wd, "A", coefs=coefs, inputs_override=inp)
    for k in before:
        assert before[k].equals(coefs[k])


def test_user_added_input_after_blank_flows_through(wd, coefs):
    """A single edit moves the path (sanity of the user workflow) and Exercise A with prefilled-equal defaults equals blank defaults."""
    base = eng.run_scenario(wd, "A", coefs=coefs).forecast["gcpi"]
    inp = wd.inputs["A"].copy()
    inp.loc[pd.Period("2023Q3", "Q"), "brent"] = 100.0
    up = eng.run_scenario(wd, "A", coefs=coefs, inputs_override=inp).forecast["gcpi"]
    assert up.iloc[0] > base.iloc[0]
    d = eng.derive_history(wd.history)
    inp2 = wd.inputs["A"].copy()
    inp2["brent"] = float(wd.history["brent"].iloc[-1])
    inp2["vu"] = float(wd.history["vu"].iloc[-1])
    same = eng.run_scenario(wd, "A", coefs=coefs, inputs_override=inp2).forecast["gcpi"]
    assert np.allclose(same.values, base.values, atol=1e-12)


# ---------------------------------------------------------------- workbook formulas agree with the engine (needs Excel)
def test_excel_formulas_match_engine(tmp_path, coefs):
    pytest.importorskip("win32com.client")
    from bb_inflation.phase5 import excel_util
    wb = load_workbook(WB)
    ws = wb["Inputs_A"]
    ws["C10"], ws["D11"], ws["E10"], ws["F11"], ws["G12"], ws["H13"], ws["J9"] = 90, 5, 6, 1.0, 1.2, 2.0, 0.5
    f = tmp_path / "t.xlsx"
    wb.save(f)
    try:
        excel_util.recalc(f)
    except Exception as e:                                         # noqa: BLE001
        pytest.skip(f"Excel automation unavailable: {e}")
    wd = eng.read_workbook(f)
    mi = eng.run_scenario(wd, "A", coefs=coefs).inputs_used
    m = load_workbook(f, data_only=True)["Model_Inputs_A"]
    hdr = {c.value: i + 1 for i, c in enumerate(m[6])}
    cols = {"brent": "Brent used", "gbrent": "Brent growth", "gE": "CPI energy growth USED (model gE)", "gscpi": "GSCPI used",
            "shortagema2": "Shortage term (2-qtr mean)", "vu": "V/U used", "gpty": "Productivity growth used",
            "magpty": "Trend productivity magpty (8-qtr mean)", "af_gcpi": "Add-factor gcpi"}
    for i, q in enumerate(mi.index):
        row = 18 + i
        assert m.cell(row, 1).value == str(q)
        for k, lab in cols.items():
            assert m.cell(row, hdr[lab]).value == pytest.approx(mi.loc[q, k], abs=1e-8), (q, k)


def test_workbook_cached_values_present_for_engine_sheets():
    wb = load_workbook(WB, data_only=True)
    assert wb["Model_Inputs_A"]["F20"].value is not None
    assert wb["Settings"]["B7"].value == "2026Q2"


# ---------------------------------------------------------------- evaluation
def test_decomposition_is_additive(p5, wd, coefs):
    q = pd.period_range("2023Q3", "2026Q2", freq="Q")
    in_a, in_b = wd.inputs["A"].loc[q[0]:q[-1]], wd.inputs["B"].loc[q[0]:q[-1]]
    shap, _ = EV.shapley_drivers(wd, coefs, in_a, in_b, "2026Q2")
    a = eng.run_scenario(wd, "A", coefs=coefs).forecast["gcpi"]
    b = eng.run_scenario(wd, "B", coefs=coefs).forecast["gcpi"]
    assert np.allclose(shap.iloc[:, :-1].sum(axis=1).values, (b - a).values, atol=1e-10)


def test_metrics_basic(p5):
    q = pd.period_range("2023Q3", "2026Q2", freq="Q")
    real = EV.realised_paths(p5.levels, q)
    m = EV.build_tables({"x": real["gcpi"] + 1.0}, real, p5.model["gcpi"].loc[:"2023Q2"])
    mm = m["metrics"].query("measure == 'CPI inflation q/q annualised'")
    assert np.allclose(mm["ME"], 1.0) and np.allclose(mm["RMSE"], 1.0) and np.allclose(mm["MAE"], 1.0)
    # y/y chaining reproduces the realised y/y when fed the realised quarterly path
    m2 = EV.build_tables({"x": real["gcpi"]}, real, p5.model["gcpi"].loc[:"2023Q2"])
    assert np.allclose(m2["errors_yoy"]["x"].values, 0.0, atol=1e-9)
