#!/usr/bin/env python
"""Phase 5 forecasting engine for Model B* (four-equation SVAR-X), driven by the Excel assumptions workbook.

Standalone: needs only numpy, pandas and openpyxl (no Macrobond connection, no estimation code, no other project module).

    python p5_forecast_engine.py --workbook Phase5_Assumptions.xlsx --scenario A          # pseudo-real-time (sheet Inputs_A)
    python p5_forecast_engine.py --workbook Phase5_Assumptions.xlsx --scenario B          # conditional      (sheet Inputs_B)
    python p5_forecast_engine.py --workbook Phase5_Assumptions.xlsx --scenario A --horizon 2029Q2   # extend (max 2029Q2)

What it does
------------
1. Reads the workbook: `History` (raw LEVELS through the 2023Q2 origin; the only data used for initial conditions), `Inputs_<A|B>`
   (the user's assumptions in intuitive units), `Settings`, `Aux_Models` (oil bridge, GSCPI AR(1)). The sheet `Realised` is read only
   when --with-realised is given and only to compute forecast errors AFTER the forecast; it never enters the simulation.
2. Reads the estimated 2023Q2 coefficients from the CSV named in Settings (or --coefs). They are never re-estimated here.
3. Converts every user assumption to the exact transformation Model B* uses (see `transform_inputs`; documented in the workbook sheet
   `Variables`): Brent level -> 400*dln -> CPI-energy growth through the bridge; CPI energy / food / productivity annualised % ->
   400*dln = 100*ln(1+x/100); GSCPI level -> 2-quarter trailing mean; V/U level as is; productivity -> 8-quarter trailing mean.
4. Anything the user leaves blank is filled by a documented default rule (see DEFAULT_RULES); every fill is written to the Engine_Log.
5. Simulates the four equations recursively from the 2023Q2 initial conditions: within a quarter
       gw -> grpe, grpf (= nominal energy/food growth - gw) -> gcpi -> catch-up -> cf10 -> cf1
   which is the exact solution of the lower-triangular structural system A0 y_t = c + sum A_i y_{t-i} + sum B_j x_{t-j}; the engine
   re-solves the system with the matrices (A0^-1) as a check and reports the difference.

Units (Model B*): gw, gcpi, grpe, grpf, gpty = 400*dln (annualised q/q, % points); cf1, cf10 = % per year (levels); vu = ratio;
shortage = GSCPI standard deviations. CPI and its components are quarter-END index values (the authors' convention).
"""
from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

LAGS = 4
Y = ["gw", "gcpi", "cf10", "cf1"]                                   # order of the structural system (as in the SVAR-X note)
XV = ["magpty", "vu", "diffcpicf", "grpe", "grpf", "shortagema2", "D2020Q2", "D2020Q3"]
MAX_HORIZON = "2029Q2"

# Model B* specification (terms permitted in each equation; anything else in the coefficient file is an error)
SPEC = {
    "gw": [f"gw_L{k}" for k in (1, 2, 3, 4)] + [f"cf1_L{k}" for k in (1, 2, 3, 4)] + ["magpty_L1"]
          + [f"vu_L{k}" for k in (1, 2, 3, 4)] + [f"diffcpicf_L{k}" for k in (1, 2, 3, 4)] + ["D2020Q2", "D2020Q3", "const"],
    "gcpi": ["magpty_L0"] + [f"gcpi_L{k}" for k in (1, 2, 3, 4)] + [f"gw_L{k}" for k in range(5)]
            + [f"grpe_L{k}" for k in range(5)] + [f"grpf_L{k}" for k in range(5)] + ["shortagema2_L0", "const"],
    "cf10": [f"cf10_L{k}" for k in (1, 2, 3, 4)] + [f"gcpi_L{k}" for k in range(5)],
    "cf1": [f"cf1_L{k}" for k in (1, 2, 3, 4)] + [f"cf10_L{k}" for k in range(5)] + [f"gcpi_L{k}" for k in range(5)],
}
HOMOGENEITY = {"gw": ("gw", "cf1"), "gcpi": ("gcpi", "gw"), "cf10": ("cf10", "gcpi"), "cf1": ("cf1", "cf10", "gcpi")}

RAW_KEYS = ["brent", "cpi", "cpi_energy", "cpi_food", "eci", "prod", "gscpi", "vu", "cf1", "cf10"]
INPUT_KEYS = ["brent", "cpie_qoq", "cpif_qoq", "gscpi", "vu", "prod_qoq", "af_gw", "af_gcpi", "af_cf10", "af_cf1"]

DEFAULT_RULES = {
    "brent": "carried forward from the previous quarter's level (history: 2023Q2 average); an entered value sticks until changed",
    "cpie_qoq": "CPI energy growth from the oil bridge (Aux_Models) applied to the Brent path",
    "cpif_qoq": "wage-linked: nominal CPI food grows with simulated wages, i.e. the relative food price is unchanged (grpf = 0)",
    "gscpi": "AR(1) mean reversion (Aux_Models) from the previous quarter's GSCPI",
    "vu": "carried forward from the previous quarter (history: 2023Q2 value)",
    "prod_qoq": "trend held: growth equals the growth 8 quarters earlier, so the 8-quarter trend (magpty) stays at its last value",
    "af": "zero",
}


# ----------------------------------------------------------------------------------------------------------------------
# Reading the workbook
# ----------------------------------------------------------------------------------------------------------------------
class WorkbookError(ValueError):
    pass


def _key_row(ws) -> int:
    for r in range(1, 40):
        if ws.cell(r, 1).value == "key":
            return r
    raise WorkbookError(f"sheet '{ws.title}': no machine-key row (cell A = 'key') found")


def read_table(wb, wbf, sheet: str, numeric: bool = True) -> pd.DataFrame:
    """Table with a key row (A='key'), then a label row and a unit row, then data. Returns a frame indexed by quarter."""
    ws, wsf = wb[sheet], wbf[sheet]
    r0 = _key_row(ws)
    keys = [ws.cell(r0, c).value for c in range(1, ws.max_column + 1)]
    rows, idx = [], []
    for r in range(r0 + 3, ws.max_row + 1):
        q = ws.cell(r, 1).value
        if q in (None, ""):
            continue
        vals = {}
        for c, k in enumerate(keys, start=1):
            if k in (None, "", "key", "quarter"):
                continue
            v = ws.cell(r, c).value
            vf = wsf.cell(r, c).value
            if v is None and isinstance(vf, str) and vf.startswith("="):
                raise WorkbookError(f"{sheet}!{ws.cell(r, c).coordinate} holds a formula without a stored value; open the workbook "
                                    f"in Excel, let it recalculate, save, and run again")
            if numeric and v is not None and not isinstance(v, (int, float)):
                raise WorkbookError(f"{sheet}!{ws.cell(r, c).coordinate} = {v!r}: a number (or an empty cell) is required")
            vals[k] = np.nan if v is None else float(v)
        rows.append(vals)
        idx.append(pd.Period(str(q).replace(" ", ""), "Q"))
    return pd.DataFrame(rows, index=pd.PeriodIndex(idx, freq="Q"))


def read_keyvalue(wb, sheet: str) -> dict:
    ws = wb[sheet]
    r0 = _key_row(ws)
    out = {}
    for r in range(r0 + 1, ws.max_row + 1):
        k, v = ws.cell(r, 1).value, ws.cell(r, 2).value
        if k not in (None, ""):
            out[str(k)] = v
    return out


@dataclass
class WorkbookData:
    path: Path
    history: pd.DataFrame           # raw levels through the origin
    inputs: dict[str, pd.DataFrame]  # 'A', 'B'
    settings: dict
    aux: dict
    realised: pd.DataFrame | None
    coef_display: pd.DataFrame | None


def read_workbook(path: str | Path) -> WorkbookData:
    path = Path(path)
    wb = load_workbook(path, data_only=True)
    wbf = load_workbook(path, data_only=False)
    hist = read_table(wb, wbf, "History")[RAW_KEYS]
    inputs = {s: read_table(wb, wbf, f"Inputs_{s}")[INPUT_KEYS] for s in ("A", "B") if f"Inputs_{s}" in wb.sheetnames}
    settings, aux = read_keyvalue(wb, "Settings"), read_keyvalue(wb, "Aux_Models")
    realised = read_table(wb, wbf, "Realised")[RAW_KEYS] if "Realised" in wb.sheetnames else None
    cd = None
    if "Coefficients" in wb.sheetnames:
        ws = wb["Coefficients"]
        r0 = _key_row(ws)
        cd = pd.DataFrame([{"eq": ws.cell(r, 2).value, "term": ws.cell(r, 3).value, "coef": ws.cell(r, 4).value}
                           for r in range(r0 + 3, ws.max_row + 1) if ws.cell(r, 2).value is not None])
    return WorkbookData(path, hist, inputs, settings, aux, realised, cd)


def read_coefficients(path: str | Path) -> dict[str, pd.Series]:
    df = pd.read_csv(path)
    coefs = {eq: g.set_index("term")["coef"].astype(float) for eq, g in df.groupby("eq")}
    validate_coefficients(coefs)
    return coefs


def validate_coefficients(coefs: dict[str, pd.Series], tol: float = 1e-6) -> None:
    """The file must contain exactly the Model B* terms (the zero restrictions) and satisfy the homogeneity restrictions."""
    if set(coefs) != set(SPEC):
        raise ValueError(f"coefficient file equations {sorted(coefs)} != {sorted(SPEC)}")
    for eq, terms in SPEC.items():
        got, want = set(coefs[eq].index), set(terms)
        if got != want:
            raise ValueError(f"equation {eq}: unexpected terms {sorted(got - want)}, missing terms {sorted(want - got)}")
        s = sum(v for t, v in coefs[eq].items() if t.rsplit("_L", 1)[0] in HOMOGENEITY[eq] and "_L" in t)
        if abs(s - 1.0) > tol:
            raise ValueError(f"equation {eq}: homogeneity restriction violated, sum = {s:.8f}")


# ----------------------------------------------------------------------------------------------------------------------
# Transformations
# ----------------------------------------------------------------------------------------------------------------------
def dln400(x: pd.Series) -> pd.Series:
    """Model scaling of a level series: 400 * dln(x)."""
    return 400.0 * np.log(x).diff()


def pct_ann_to_model(x: float) -> float:
    """Annualised compounded % change x (e.g. 6.0) -> the model's 400*dln of the quarterly ratio: 400*ln((1+x/100)^(1/4)) = 100*ln(1+x/100)."""
    if not np.isfinite(x) or x <= -100.0:
        raise ValueError(f"annualised % change {x!r} is not admissible (must be > -100)")
    return 100.0 * np.log1p(x / 100.0)


def model_to_pct_ann(g: float) -> float:
    """Inverse of pct_ann_to_model: 400*dln -> annualised compounded % change."""
    return 100.0 * (np.exp(g / 100.0) - 1.0)


def derive_history(lv: pd.DataFrame) -> pd.DataFrame:
    """Model B* variables from raw levels, exactly as in the estimation data: growth rates are 400*dln of quarter-end levels, relative
    prices are 400*dln(price/ECI), trend productivity is the 8-quarter trailing mean of productivity growth, the shortage term is the
    2-quarter trailing mean of the GSCPI, the catch-up term is the 4-quarter mean of CPI inflation less 1-year expectations 4 quarters back."""
    d = pd.DataFrame(index=lv.index)
    d["gcpi"] = dln400(lv["cpi"])
    d["gw"] = dln400(lv["eci"])
    d["gE"] = dln400(lv["cpi_energy"])
    d["gF"] = dln400(lv["cpi_food"])
    d["grpe"] = dln400(lv["cpi_energy"] / lv["eci"])
    d["grpf"] = dln400(lv["cpi_food"] / lv["eci"])
    d["gpty"] = dln400(lv["prod"])
    d["magpty"] = d["gpty"].rolling(8).mean()
    d["shortage"] = lv["gscpi"]
    d["shortagema2"] = lv["gscpi"].rolling(2).mean()
    d["vu"], d["cf1"], d["cf10"] = lv["vu"], lv["cf1"], lv["cf10"]
    d["diffcpicf"] = d["gcpi"].rolling(4).mean() - d["cf1"].shift(4)
    return d


@dataclass
class Transformed:
    table: pd.DataFrame                 # one row per forecast quarter
    log: list[dict] = field(default_factory=list)


def transform_inputs(lv: pd.DataFrame, d: pd.DataFrame, inp: pd.DataFrame, aux: dict, quarters: pd.PeriodIndex) -> Transformed:
    """Convert the user's assumptions (intuitive units, blank = default rule) to the Model B* inputs for the forecast quarters.

    Historical-to-forecast transition: every recursion starts from the last history quarter (the origin): the Brent growth of the first
    forecast quarter uses the origin-quarter Brent average, GSCPI's 2-quarter mean uses the origin GSCPI, the productivity trend uses the
    last 8 history growth rates.
    """
    for k in ("bridge_const", "bridge_b0", "gscpi_ar_const", "gscpi_ar_rho"):
        if k not in aux or not isinstance(aux[k], (int, float)):
            raise WorkbookError(f"Aux_Models is missing a numeric value for '{k}'")
    log: list[dict] = []
    brent_prev = float(lv["brent"].iloc[-1])
    gscpi_prev = float(lv["gscpi"].iloc[-1])
    vu_prev = float(lv["vu"].iloc[-1])
    gpty_path = list(d["gpty"].iloc[-8:].astype(float))
    rows = []
    for q in quarters:
        u = inp.loc[q] if q in inp.index else pd.Series(np.nan, index=INPUT_KEYS)
        r: dict = {"quarter": q}

        def note(var, rule, detail=""):
            log.append({"quarter": str(q), "variable": var, "rule": rule, "detail": detail})

        # --- Brent (level, US$/bbl) -> 400*dln -> CPI-energy growth through the bridge
        if np.isfinite(u["brent"]):
            if u["brent"] <= 0:
                raise ValueError(f"{q}: Brent must be positive, got {u['brent']}")
            brent, r["brent_src"] = float(u["brent"]), "user"
        else:
            brent, r["brent_src"] = brent_prev, "default"
            note("brent", "default", DEFAULT_RULES["brent"])
        r["brent"] = brent
        r["gbrent"] = 400.0 * np.log(brent / brent_prev)
        r["gE_bridge"] = float(aux["bridge_const"]) + float(aux["bridge_b0"]) * r["gbrent"]
        brent_prev = brent
        # --- CPI energy: user override (annualised % q/q) else bridge
        if np.isfinite(u["cpie_qoq"]):
            r["gE"], r["gE_src"] = pct_ann_to_model(u["cpie_qoq"]), "user (CPI energy override)"
        else:
            r["gE"] = r["gE_bridge"]
            r["gE_src"] = "oil bridge, Brent entered by user" if r["brent_src"] == "user" else "oil bridge, Brent default (carried forward)"
            if r["brent_src"] != "user":
                note("cpi_energy", "default", DEFAULT_RULES["cpie_qoq"])
        # --- CPI food: user (annualised % q/q) else wage-linked (NaN here; resolved in the simulation)
        if np.isfinite(u["cpif_qoq"]):
            r["gF"], r["gF_src"] = pct_ann_to_model(u["cpif_qoq"]), "user"
        else:
            r["gF"], r["gF_src"] = np.nan, "default: wage-linked (relative food price unchanged)"
            note("cpi_food", "default", DEFAULT_RULES["cpif_qoq"])
        # --- GSCPI (level) -> 2-quarter trailing mean
        if np.isfinite(u["gscpi"]):
            g, r["gscpi_src"] = float(u["gscpi"]), "user"
        else:
            g, r["gscpi_src"] = float(aux["gscpi_ar_const"]) + float(aux["gscpi_ar_rho"]) * gscpi_prev, "default: AR(1)"
            note("gscpi", "default", DEFAULT_RULES["gscpi"])
        r["gscpi"] = g
        r["shortagema2"] = 0.5 * (g + gscpi_prev)
        gscpi_prev = g
        # --- V/U (level)
        if np.isfinite(u["vu"]):
            if u["vu"] < 0:
                raise ValueError(f"{q}: V/U cannot be negative, got {u['vu']}")
            vu, r["vu_src"] = float(u["vu"]), "user"
        else:
            vu, r["vu_src"] = vu_prev, "default: carried forward"
            note("vu", "default", DEFAULT_RULES["vu"])
        r["vu"] = vu
        vu_prev = vu
        # --- productivity (annualised % q/q) -> 400*dln -> 8-quarter trailing mean
        if np.isfinite(u["prod_qoq"]):
            gp, r["prod_src"] = pct_ann_to_model(u["prod_qoq"]), "user"
        else:
            gp, r["prod_src"] = gpty_path[-8], "default: trend held"
            note("productivity", "default", DEFAULT_RULES["prod_qoq"])
        gpty_path.append(gp)
        r["gpty"] = gp
        r["magpty"] = float(np.mean(gpty_path[-8:]))
        # --- add-factors (pp, annualised), added to the equation
        for e in ("gw", "gcpi", "cf10", "cf1"):
            v = u[f"af_{e}"]
            if np.isfinite(v):
                r[f"af_{e}"], r[f"af_{e}_src"] = float(v), "user"
            else:
                r[f"af_{e}"], r[f"af_{e}_src"] = 0.0, "default: zero"
        rows.append(r)
    return Transformed(pd.DataFrame(rows).set_index("quarter"), log)


# ----------------------------------------------------------------------------------------------------------------------
# Simulation
# ----------------------------------------------------------------------------------------------------------------------
def _terms(coefs: dict[str, pd.Series]) -> dict[str, list[tuple[str, int, float]]]:
    out = {}
    for eq, c in coefs.items():
        tl = []
        for name, b in c.items():
            if "_L" in name:
                var, lag = name.rsplit("_L", 1)
                tl.append((var, int(lag), float(b)))
        out[eq] = tl
    return out


def simulate(coefs: dict[str, pd.Series], d: pd.DataFrame, mi: pd.DataFrame,
             grpe_override: np.ndarray | None = None, grpf_override: np.ndarray | None = None,
             const_override: dict[str, float] | None = None, extra_add: dict[str, np.ndarray] | None = None) -> pd.DataFrame:
    """Recursive simulation from the last LAGS history rows of `d`. Nothing after the origin enters except the transformed inputs `mi`.

    grpe_override / grpf_override (arrays over the forecast quarters, NaN = derive): used only for the 'realised relative prices'
    diagnostic variant. const_override / extra_add: diagnostics (decompositions)."""
    H = len(mi)
    T = LAGS + H
    h = d.iloc[-LAGS:]
    s = {k: np.full(T, np.nan) for k in ("gw", "gcpi", "cf1", "cf10", "diffcpicf", "grpe", "grpf", "vu", "magpty", "shortagema2")}
    for k in s:
        s[k][:LAGS] = h[k].values.astype(float)
    s["vu"][LAGS:] = mi["vu"].values
    s["magpty"][LAGS:] = mi["magpty"].values
    s["shortagema2"][LAGS:] = mi["shortagema2"].values
    dummies = {"D2020Q2": 0.0, "D2020Q3": 0.0}                    # no pandemic quarters inside a forecast horizon
    terms = _terms(coefs)
    const_override = const_override or {}
    extra_add = extra_add or {}

    def val(eq: str, t: int) -> float:
        v = sum(b * s[var][t - lag] for var, lag, b in terms[eq])
        c = coefs[eq]
        for dn, dv in dummies.items():
            if dn in c.index:
                v += c[dn] * dv
        v += const_override.get(eq, c["const"] if "const" in c.index else 0.0)
        v += mi[f"af_{eq}"].values[t - LAGS]
        if eq in extra_add:
            v += extra_add[eq][t - LAGS]
        return v

    for t in range(LAGS, T):
        j = t - LAGS
        s["gw"][t] = val("gw", t)
        gE, gF = mi["gE"].values[j], mi["gF"].values[j]
        s["grpe"][t] = grpe_override[j] if (grpe_override is not None and np.isfinite(grpe_override[j])) else gE - s["gw"][t]
        if grpf_override is not None and np.isfinite(grpf_override[j]):
            s["grpf"][t] = grpf_override[j]
        else:
            s["grpf"][t] = (gF - s["gw"][t]) if np.isfinite(gF) else 0.0      # NaN = wage-linked food: relative price unchanged
        s["gcpi"][t] = val("gcpi", t)
        s["diffcpicf"][t] = 0.25 * (s["gcpi"][t] + s["gcpi"][t - 1] + s["gcpi"][t - 2] + s["gcpi"][t - 3]) - s["cf1"][t - 4]
        s["cf10"][t] = val("cf10", t)
        s["cf1"][t] = val("cf1", t)
    out = pd.DataFrame({k: s[k][LAGS:] for k in ("gw", "gcpi", "cf10", "cf1", "diffcpicf", "grpe", "grpf")}, index=mi.index)
    return out


def build_matrices(coefs: dict[str, pd.Series]):
    """Structural matrices of  A0 y_t = c + sum_i A_i y_{t-i} + sum_j B_j x_{t-j} + e_t  (rows/cols ordered Y, XV)."""
    A = {i: np.zeros((4, 4)) for i in range(LAGS + 1)}
    B = {j: np.zeros((4, len(XV))) for j in range(LAGS + 1)}
    c = np.zeros(4)
    for eq, cs in coefs.items():
        r = Y.index(eq)
        for name, b in cs.items():
            if name == "const":
                c[r] = b
            elif name in ("D2020Q2", "D2020Q3"):
                B[0][r, XV.index(name)] = b
            else:
                var, lag = name.rsplit("_L", 1)
                lag = int(lag)
                if var in Y:
                    A[lag][r, Y.index(var)] += (-b if lag == 0 else b)      # A0 holds minus the contemporaneous coefficients
                else:
                    B[lag][r, XV.index(var)] = b
    A[0] += np.eye(4)
    return A, B, c


def matrix_form_check(coefs: dict[str, pd.Series], d: pd.DataFrame, mi: pd.DataFrame, path: pd.DataFrame) -> float:
    """Re-solve y_t = A0^{-1}(c + sum A_i y_{t-i} + sum B_j x_{t-j} + af) with the structural matrices, taking the exogenous drivers
    (incl. the relative prices of the main run) as given; return the maximum absolute difference to the sequential solution."""
    A, B, c = build_matrices(coefs)
    A0inv = np.linalg.inv(A[0])
    H = len(mi)
    h = d.iloc[-LAGS:]
    yv = np.full((LAGS + H, 4), np.nan)
    xv = np.full((LAGS + H, len(XV)), np.nan)
    for i, v in enumerate(Y):
        yv[:LAGS, i] = h[v].values
    for i, v in enumerate(XV):
        if v in h:
            xv[:LAGS, i] = h[v].values
    xv[LAGS:, XV.index("magpty")] = mi["magpty"].values
    xv[LAGS:, XV.index("vu")] = mi["vu"].values
    xv[LAGS:, XV.index("shortagema2")] = mi["shortagema2"].values
    xv[LAGS:, XV.index("grpe")] = path["grpe"].values
    xv[LAGS:, XV.index("grpf")] = path["grpf"].values
    xv[LAGS:, XV.index("diffcpicf")] = 0.0                # filled in as each quarter is solved (no contemporaneous coefficient)
    xv[:, XV.index("D2020Q2")] = 0.0                      # no pandemic dummies in the lag window or the horizon
    xv[:, XV.index("D2020Q3")] = 0.0
    cf1_i, gcpi_i = Y.index("cf1"), Y.index("gcpi")
    for t in range(LAGS, LAGS + H):
        rhs = c + sum(A[i] @ yv[t - i] for i in range(1, LAGS + 1)) + sum(B[j] @ xv[t - j] for j in range(LAGS + 1))
        rhs = rhs + np.array([mi[f"af_{e}"].values[t - LAGS] for e in Y])
        yv[t] = A0inv @ rhs
        xv[t, XV.index("diffcpicf")] = 0.25 * (yv[t, gcpi_i] + yv[t - 1, gcpi_i] + yv[t - 2, gcpi_i] + yv[t - 3, gcpi_i]) - yv[t - 4, cf1_i]
    return float(np.abs(yv[LAGS:] - path[Y].values).max())


# ----------------------------------------------------------------------------------------------------------------------
# Running a scenario
# ----------------------------------------------------------------------------------------------------------------------
@dataclass
class Result:
    scenario: str
    forecast: pd.DataFrame
    inputs_used: pd.DataFrame
    log: pd.DataFrame
    checks: dict
    horizon: str
    meta: dict


def levels_paths(lv: pd.DataFrame, mi: pd.DataFrame, path: pd.DataFrame) -> pd.DataFrame:
    """Index-level paths implied by the forecast (quarter-end CPI, ECI, CPI energy, CPI food) and CPI y/y."""
    cpi, eci, ce, cf = (float(lv[k].iloc[-1]) for k in ("cpi", "eci", "cpi_energy", "cpi_food"))
    hist = lv[["cpi", "cpi_energy", "cpi_food", "eci"]].copy()
    rows = []
    for q in path.index:
        cpi *= np.exp(path.at[q, "gcpi"] / 400.0)
        eci *= np.exp(path.at[q, "gw"] / 400.0)
        ce *= np.exp(mi.at[q, "gE"] / 400.0)
        gF = mi.at[q, "gF"]
        cf *= np.exp((path.at[q, "gw"] if not np.isfinite(gF) else gF) / 400.0)
        rows.append({"quarter": q, "cpi": cpi, "cpi_energy": ce, "cpi_food": cf, "eci": eci})
    fut = pd.DataFrame(rows).set_index("quarter")
    allv = pd.concat([hist, fut])
    out = fut.copy()
    out["cpi_yoy"] = 100.0 * (allv["cpi"] / allv["cpi"].shift(4) - 1.0).loc[fut.index]
    out["cpi_energy_yoy"] = 100.0 * (allv["cpi_energy"] / allv["cpi_energy"].shift(4) - 1.0).loc[fut.index]
    out["cpi_food_yoy"] = 100.0 * (allv["cpi_food"] / allv["cpi_food"].shift(4) - 1.0).loc[fut.index]
    out["eci_yoy"] = 100.0 * (allv["eci"] / allv["eci"].shift(4) - 1.0).loc[fut.index]
    return out


def run_scenario(workbook: str | Path | WorkbookData, scenario: str, coef_path: str | Path | None = None, horizon: str | None = None,
                 inputs_override: pd.DataFrame | None = None, coefs: dict[str, pd.Series] | None = None,
                 grpe_override: np.ndarray | None = None, grpf_override: np.ndarray | None = None,
                 extra_add: dict[str, np.ndarray] | None = None, with_realised: bool = False) -> Result:
    wd = workbook if isinstance(workbook, WorkbookData) else read_workbook(workbook)
    scenario = scenario.upper()
    if scenario not in wd.inputs:
        raise KeyError(f"workbook has no sheet Inputs_{scenario}")
    horizon = str(horizon or wd.settings.get("forecast_end") or "2026Q2").replace(" ", "")
    if pd.Period(horizon, "Q") > pd.Period(MAX_HORIZON, "Q"):
        raise ValueError(f"forecast_end {horizon} is beyond the workbook's last input quarter {MAX_HORIZON}")
    origin = wd.history.index[-1]
    if str(origin) != str(wd.settings.get("origin", origin)).replace(" ", ""):
        raise WorkbookError(f"History ends {origin} but Settings says origin = {wd.settings.get('origin')}")
    if coefs is None:
        cp = Path(coef_path) if coef_path else (wd.path.parent / str(wd.settings["coef_file"]))
        coefs = read_coefficients(cp)
    else:
        validate_coefficients(coefs)
    quarters = pd.period_range(origin + 1, pd.Period(horizon, "Q"), freq="Q")
    if len(quarters) < 1:
        raise ValueError("forecast_end must be after the origin")
    lv = wd.history
    d = derive_history(lv)
    if d[["gw", "gcpi", "cf1", "cf10", "diffcpicf", "grpe", "grpf", "magpty"]].iloc[-LAGS:].isna().any().any():
        raise WorkbookError("History is too short to initialise the lags (need 4 + 8 quarters of productivity growth + 4 for the catch-up)")
    inp = wd.inputs[scenario] if inputs_override is None else inputs_override
    tr = transform_inputs(lv, d, inp, wd.aux, quarters)
    path = simulate(coefs, d, tr.table, grpe_override, grpf_override, extra_add=extra_add)
    maxdiff = matrix_form_check(coefs, d, tr.table, path) if extra_add is None else np.nan
    paths = levels_paths(lv, tr.table, path)
    fc = path.join(paths)
    fc["gcpi_yoy_4q_mean"] = pd.concat([d["gcpi"], fc["gcpi"]]).rolling(4).mean().loc[fc.index]
    if with_realised and wd.realised is not None:
        rl = wd.realised
        full = pd.concat([lv, rl.loc[rl.index > origin]])
        g = 400.0 * np.log(full["cpi"]).diff()
        yy = 100.0 * (full["cpi"] / full["cpi"].shift(4) - 1.0)
        fc["gcpi_realised"] = g.reindex(fc.index)
        fc["cpi_yoy_realised"] = yy.reindex(fc.index)
        fc["error_gcpi"] = fc["gcpi"] - fc["gcpi_realised"]
        fc["error_cpi_yoy"] = fc["cpi_yoy"] - fc["cpi_yoy_realised"]
    disp_diff = np.nan
    if wd.coef_display is not None and len(wd.coef_display):
        ref = pd.concat({eq: c for eq, c in coefs.items()}, names=["eq", "term"]).rename("ref")
        disp = wd.coef_display.set_index(["eq", "term"])["coef"].astype(float)
        j = pd.concat([ref, disp.rename("disp")], axis=1)
        disp_diff = float((j["ref"] - j["disp"]).abs().max()) if len(j) else np.nan
        if not (disp_diff <= 1e-9):
            print(f"WARNING: coefficients used differ from the display sheet 'Coefficients' (max abs diff {disp_diff:.3g}; "
                  f"different file or different vintage?)", file=sys.stderr)
    checks = {"matrix_form_max_abs_diff": maxdiff, "coef_display_vs_used_max_abs_diff": disp_diff, "origin": str(origin), "horizon": horizon, "n_quarters": len(quarters),
              "coefficient_file": str(coef_path or wd.settings.get("coef_file")), "realised_used_in_simulation": False}
    meta = {"workbook": str(wd.path)}
    return Result(scenario, fc, tr.table, pd.DataFrame(tr.log), checks, horizon, meta)


def write_result(res: Result, out_dir: str | Path, tag: str | None = None) -> list[Path]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    tag = tag or f"forecast_{res.scenario}"
    f = out / f"{tag}.csv"
    res.forecast.reset_index().assign(quarter=lambda x: x["quarter"].astype(str)).to_csv(f, index=False)
    x = out / f"{tag}.xlsx"
    with pd.ExcelWriter(x, engine="openpyxl") as w:
        res.forecast.reset_index().assign(quarter=lambda t: t["quarter"].astype(str)).to_excel(w, sheet_name="Forecast", index=False)
        res.inputs_used.reset_index().assign(quarter=lambda t: t["quarter"].astype(str)).to_excel(w, sheet_name="Model_Inputs_Engine", index=False)
        (res.log if len(res.log) else pd.DataFrame(columns=["quarter", "variable", "rule", "detail"])).to_excel(w, sheet_name="Engine_Log", index=False)
        pd.Series(res.checks).rename("value").to_frame().reset_index().rename(columns={"index": "check"}).to_excel(w, sheet_name="Checks", index=False)
    return [f, x]


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1], formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--scenario", default="A", help="A (sheet Inputs_A) or B (sheet Inputs_B)")
    ap.add_argument("--coefs", default=None, help="coefficient CSV (default: the file named in Settings, relative to the workbook)")
    ap.add_argument("--horizon", default=None, help="last forecast quarter, e.g. 2026Q2 (default: Settings!forecast_end; max 2029Q2)")
    ap.add_argument("--out-dir", default=None, help="default: <workbook folder>/output")
    ap.add_argument("--with-realised", action="store_true", help="append realised CPI inflation and errors (reporting only)")
    a = ap.parse_args(argv)
    res = run_scenario(a.workbook, a.scenario, a.coefs, a.horizon, with_realised=a.with_realised)
    out = Path(a.out_dir) if a.out_dir else Path(a.workbook).resolve().parent / "output"
    files = write_result(res, out)
    pd.set_option("display.width", 200)
    cols = ["gcpi", "cpi_yoy", "gw", "cf1", "cf10"] + (["gcpi_realised", "error_gcpi"] if a.with_realised else [])
    print(res.forecast[cols].round(2).to_string())
    print(f"\nscenario {res.scenario}: {res.checks['n_quarters']} quarters from {res.checks['origin']}; matrix-form check max|diff| = "
          f"{res.checks['matrix_form_max_abs_diff']:.2e}; {len(res.log)} default-rule fills (see Engine_Log)")
    for f in files:
        print("wrote", f)
    return 0


if __name__ == "__main__":
    sys.exit(main())
